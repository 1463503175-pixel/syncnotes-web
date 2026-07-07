import html
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree


MAX_CHARS = 16000


def compact(text):
    text = re.sub(r"\s+", " ", text or "").strip()
    return text[:MAX_CHARS]


def read_text(path):
    data = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    chunks = []
    for page in reader.pages[:30]:
        chunks.append(page.extract_text() or "")
    return "\n".join(chunks)


def extract_docx(path):
    from docx import Document

    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables[:20]:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def slide_order(name):
    match = re.search(r"slide(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def extract_pptx(path):
    chunks = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            [n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
            key=slide_order,
        )
        for name in names[:80]:
            xml = zf.read(name)
            try:
                root = ElementTree.fromstring(xml)
            except ElementTree.ParseError:
                continue
            texts = []
            for node in root.iter():
                if node.tag.endswith("}t") and node.text:
                    texts.append(node.text)
            if texts:
                chunks.append(" ".join(texts))
    return "\n".join(chunks)


def extract_xlsx(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for sheet in workbook.worksheets[:8]:
        chunks.append(f"[{sheet.title}]")
        for row in sheet.iter_rows(max_row=80, max_col=20, values_only=True):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if values:
                chunks.append(" | ".join(values))
    workbook.close()
    return "\n".join(chunks)


def extract_htmlish(path):
    text = read_text(path)
    text = re.sub(r"<script[\s\S]*?</script>", " ", text, flags=re.I)
    text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    return html.unescape(text)


def main():
    if len(sys.argv) < 3:
        raise SystemExit("usage: extract_attachment_text.py <path> <kind>")
    path = sys.argv[1]
    kind = (sys.argv[2] or "").lower()
    suffix = Path(path).suffix.lower()

    try:
        if kind == "pdf" or suffix == ".pdf":
            text = extract_pdf(path)
        elif kind == "word" or suffix == ".docx":
            text = extract_docx(path)
        elif kind == "ppt" or suffix == ".pptx":
            text = extract_pptx(path)
        elif kind == "sheet" or suffix == ".xlsx":
            text = extract_xlsx(path)
        elif suffix in (".html", ".htm", ".xml"):
            text = extract_htmlish(path)
        else:
            text = read_text(path)
        print(json.dumps({"ok": True, "text": compact(text)}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)[:300]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
